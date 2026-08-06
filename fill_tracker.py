#!/usr/bin/env python3
"""Fill the Actions Tracker spreadsheet from MOE-format minutes of meeting (.docx).

Deterministic: same inputs -> same output. No LLM.

Minutes template (per-file):
  - header table  : عنوان الاجتماع / التاريخ / الموقع ...
  - attendees table : الجهة | الاسم
  - "المطلوب" table : ملاحظات | الوقت المتوقع/المقترح للانتهاء | المسؤول | المهمة

Each action row becomes a row in the tracker's 'سجل المهام' sheet, mapping:
  المهمة  -> المهمة
  المسؤول -> المسؤول
  الوقت المتوقع للانتهاء -> تاريخ التسليم المستهدف + تاريخ المستهدف وفق المحضر
  ملاحظات -> ملاحظات
  meeting title + date -> الاجتماع / المصدر
  (new) -> الحالة = لم تبدأ,  آخر تحديث = today

Usage:
  python fill_tracker.py --tracker "Actions Tracker.xlsx" \
      --minutes "3.3 Minutes of Meeting" [--out filled.xlsx] [--status "لم تبدأ"]
"""

import argparse
import datetime
import os
import sys
import zipfile

import docx
from openpyxl import load_workbook
from docx.oxml.ns import qn

DEFAULT_STATUS = "لم تبدأ"              # status for newly added actions
TRACKER_SHEET = "سجل المهام"            # sheet to fill
DUE_MARKERS = ("انتهاء",)               # time field: الوقت المتوقع/المقترح للانتهاء


# --------------------------------------------------------------------------
# minutes .docx parsing
# --------------------------------------------------------------------------
def _blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            pass
        elif child.tag == qn("w:tbl"):
            from docx.table import Table
            yield Table(child, document)


def _rows(table):
    return [[c.text.strip() for c in row.cells] for row in table.rows]


def parse_minutes(path):
    """Return list of action dicts parsed from one minutes .docx."""
    doc = docx.Document(path)
    tables = list(_blocks(doc))

    title = date = ""
    if tables:
        for row in _rows(tables[0]):
            if "عنوان الاجتماع" in row:
                title = next((c for c in row if c and c != "عنوان الاجتماع"), "")
            if "التاريخ" in row:
                i = row.index("التاريخ")
                if i - 1 >= 0:
                    date = row[i - 1]

    team_names = []
    for tb in tables:
        data = _rows(tb)
        if len(tb.columns) == 2 and data and "الاسم" in data[0]:
            for ent, name in data[1:]:
                if "هوّز" in ent and name:
                    team_names.append(name)

    actions = []
    for tb in tables:
        data = _rows(tb)
        # locate actions header: a row containing both المهمة and المسؤول
        hdr = next((i for i, r in enumerate(data)
                    if "المهمة" in r and "المسؤول" in r), None)
        if hdr is None:
            continue
        head = data[hdr]
        ti = next((i for i, c in enumerate(head) if "المهمة" in c), None)
        ri = next((i for i, c in enumerate(head) if "المسؤول" in c), None)
        di = next((i for i, c in enumerate(head) if any(m in c for m in DUE_MARKERS)), None)
        ni = next((i for i, c in enumerate(head) if "ملاحظات" in c), None)

        for row in data[hdr + 1:]:
            if not any(row):
                continue
            task = row[ti] if ti is not None and ti < len(row) else ""
            if not task.strip():
                continue
            resp = row[ri] if ri is not None and ri < len(row) else ""
            due = row[di] if di is not None and di < len(row) else ""
            notes = row[ni] if ni is not None and ni < len(row) else ""
            actions.append({
                "task": task,
                "responsible": resp,
                "due": due,
                "notes": notes,
                "team_owner": _match_owner(team_names, task, notes, resp),
                "source": f"{title} — {date}" if (title and date) else title or date,
            })
    return actions


def _match_owner(team_names, task, notes, responsible):
    """Deterministic best-effort: find a is-hawaz member name in the text."""
    hay = f"{task} {notes} {responsible}"
    for name in team_names:
        first = name.split()[0] if name.split() else name
        if first and first in hay:
            return name
    return ""


# --------------------------------------------------------------------------
# fill tracker xlsx
# --------------------------------------------------------------------------
HEADER_MAP = {
    "المسؤول من الفريق": "team_owner",
    "المهمة": "task",
    "المسؤول": "responsible",
    "الحالة": "_status",
    "تاريخ التسليم المستهدف": "due",
    "الاجتماع / المصدر": "source",
    "ملاحظات": "notes",
    "تاريخ المستهدف وفق المحضر": "due",
    "آخر تحديث": "_updated",
}


def find_header_row(ws):
    """Return (row_idx, {header_label: col_idx}, num_col) of the actions header row.

    num_col is the column of the '#' / 'رقم' sequence field, or None.
    """
    for r in range(1, min(ws.max_row, 15) + 1):
        labels = {}
        num_col = None
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                v = v.strip().rstrip()
                if v in HEADER_MAP:
                    labels[v] = c
                elif v in ("#", "رقم", "رقم المهمة"):
                    num_col = c
        if "المهمة" in labels and "المسؤول" in labels:
            return r, labels, num_col
    return None, None, None


def build_row(action, seq, status, today):
    """Map an action onto a full tracker row keyed by header label."""
    return {
        "المسؤول من الفريق": action["team_owner"],
        "المهمة": action["task"],
        "المسؤول": action["responsible"],
        "الحالة": status,
        "تاريخ التسليم المستهدف": action["due"],
        "الاجتماع / المصدر": action["source"],
        "ملاحظات": action["notes"],
        "تاريخ المستهدف وفق المحضر": action["due"],
        "آخر تحديث": today,
    }


def fill_tracker(tracker_path, actions, status=DEFAULT_STATUS, out=None):
    wb = load_workbook(tracker_path)
    ws = wb[TRACKER_SHEET]
    hdr_row, hdr, num_col = find_header_row(ws)
    if hdr_row is None:
        raise RuntimeError(f"Could not find actions header row in sheet '{TRACKER_SHEET}'.")
    if num_col is None:
        num_col = max(1, hdr["المهمة"] - 2)  # fallback: assume '#' is two cols left

    # find first empty row below the header (rows with no task text)
    start = hdr_row + 1
    nxt = start
    while nxt <= ws.max_row:
        val = ws.cell(row=nxt, column=hdr["المهمة"]).value
        if val is None or not str(val).strip():
            break
        nxt += 1
    if nxt > ws.max_row:
        nxt = ws.max_row + 1

    seq = nxt - hdr_row          # sequence # for the first new row
    today = datetime.date.today().strftime("%d/%m/%Y")
    written = 0
    for action in actions:
        row = build_row(action, seq, status, today)
        for label, col in hdr.items():
            ws.cell(row=nxt, column=col).value = row[label]
        # write the sequence number into the '#' column
        ws.cell(row=nxt, column=num_col).value = seq
        nxt += 1
        seq += 1
        written += 1

    dest = out or tracker_path
    wb.save(dest)
    return written, dest


# --------------------------------------------------------------------------
def collect_minutes(minutes_arg):
    """Return a list of .docx paths from a folder or a zip archive."""
    files = []
    if os.path.isdir(minutes_arg):
        for root, _, names in os.walk(minutes_arg):
            for n in names:
                if n.lower().endswith(".docx"):
                    files.append(os.path.join(root, n))
    elif zipfile.is_zipfile(minutes_arg):
        z = zipfile.ZipFile(minutes_arg)
        import tempfile
        tmp = tempfile.mkdtemp(prefix="minutes_")
        for name in z.namelist():
            if name.lower().endswith(".docx"):
                z.extract(name, tmp)
        for root, _, names in os.walk(tmp):
            for n in names:
                if n.lower().endswith(".docx"):
                    files.append(os.path.join(root, n))
    else:
        raise SystemExit(f"Not a folder or zip: {minutes_arg}")
    return sorted(files)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tracker", required=True, help="Path to Actions Tracker .xlsx")
    ap.add_argument("--minutes", required=True,
                    help="Folder or .zip containing minutes .docx files")
    ap.add_argument("--out", default=None, help="Output .xlsx (default: overwrite tracker)")
    ap.add_argument("--status", default=DEFAULT_STATUS,
                    help="Status to set on new rows (default: 'لم تبدأ')")
    args = ap.parse_args(argv)

    minutes_files = collect_minutes(args.minutes)
    if not minutes_files:
        sys.exit("No .docx minutes found.")

    all_actions = []
    for f in minutes_files:
        all_actions.extend(parse_minutes(f))

    if not all_actions:
        sys.exit("No action items found in any minutes file.")

    written, dest = fill_tracker(args.tracker, all_actions,
                                 status=args.status, out=args.out)
    print(f"Parsed {len(minutes_files)} minutes files; "
          f"wrote {written} action rows -> {dest}")


if __name__ == "__main__":
    main()
